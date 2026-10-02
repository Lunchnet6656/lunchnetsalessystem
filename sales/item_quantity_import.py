"""持参数xlsx（振分表マクロ「データ変換」の出力）の読み取り・検証・保存。

画面アップロード（upload_view）と自動送信API（api_item_quantity）で共用する。
形式：シート名=対象日(yyyymmdd)、1行目C列〜=店名、2行目〜 A=対象wk(yyyymmdd)/B=商品No/C〜=数量。

振分表マクロは空白セルを詰めて書き出すため、空白が1つあるとそれ以降の数字が隣の店へズレる。
ズレた数字を黙って登録しないよう、「1つでもおかしければ1件も保存しない」。
"""
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import jpholiday
import openpyxl
from django.conf import settings
from django.db import transaction

from sales.models import Holiday, ItemQuantity, Product, SalesLocation

MAX_PRODUCT_ROW = 12  # 2〜12行目＝商品No.1〜11（11=大盛りご飯）


@dataclass
class ParsedItemQuantity:
    target_date: date | None = None
    target_week: date | None = None
    # (product, sales_location, quantity)
    entries: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    # 止めずに知らせるだけのもの（アプリで管理しない店の読み飛ばし等）
    warnings: list = field(default_factory=list)


def _parse_yyyymmdd(value):
    try:
        return datetime.strptime(str(value).strip(), "%Y%m%d").date()
    except (TypeError, ValueError):
        return None


def parse_item_quantity_workbook(fileobj) -> ParsedItemQuantity:
    result = ParsedItemQuantity()
    try:
        sheet = openpyxl.load_workbook(fileobj, data_only=True).active
    except Exception:
        result.errors.append("Excelファイルとして読み込めませんでした")
        return result

    result.target_date = _parse_yyyymmdd(sheet.title)
    if result.target_date is None:
        result.errors.append(f"シート名「{sheet.title}」が日付(yyyymmdd)ではありません")
    result.target_week = _parse_yyyymmdd(sheet.cell(row=2, column=1).value)
    if result.target_week is None:
        result.errors.append("A2の対象wkが日付(yyyymmdd)ではありません")
    if result.errors:
        return result

    locations = {loc.name: loc for loc in SalesLocation.objects.all()}
    headers = []  # (列番号, 店名, SalesLocation or None)
    for col in range(3, sheet.max_column + 1):
        name = sheet.cell(row=1, column=col).value
        if name is None or str(name).strip() == "":
            continue
        name = str(name).strip()
        headers.append((col, name, locations.get(name)))
    if not headers:
        result.errors.append("1行目に店名がありません")
        return result

    # 振分表にはあるがアプリで持参数を管理しない店（設定で明示）。それ以外の不一致は止める
    ignored = set(getattr(settings, "ITEM_QUANTITY_IGNORED_LOCATIONS", []))
    unknown = [name for _, name, loc in headers if loc is None and name not in ignored]
    if unknown:
        result.errors.append("アプリに登録されていない店名：" + "、".join(unknown))
    skipped = [name for _, name, loc in headers if loc is None and name in ignored]
    if skipped:
        result.warnings.append("アプリで管理しない店として読み飛ばし：" + "、".join(skipped))
    seen = set()
    for _, name, _ in headers:
        if name in seen:
            result.errors.append(f"店名「{name}」が2回出てきます")
        seen.add(name)

    products = {
        p.no: p for p in Product.objects.filter(week=result.target_week)
    }
    for row in range(2, MAX_PRODUCT_ROW + 1):
        product_no = sheet.cell(row=row, column=2).value
        if product_no is None:
            continue
        product = products.get(product_no)
        if product is None:
            result.errors.append(
                f"{row}行目：対象wk {result.target_week:%Y%m%d} のメニューNo.{product_no} がアプリにありません"
            )
            continue
        for col, name, loc in headers:
            value = sheet.cell(row=row, column=col).value
            if value is None or str(value).strip() == "":
                result.errors.append(
                    f"No.{product_no}（{product.name}）の「{name}」が空白です（列ズレの可能性）"
                )
                continue
            try:
                as_float = float(value)
            except (TypeError, ValueError):
                result.errors.append(f"No.{product_no}の「{name}」が数字ではありません：{value}")
                continue
            if as_float < 0 or not as_float.is_integer():
                result.errors.append(f"No.{product_no}の「{name}」が正しい個数ではありません：{value}")
                continue
            quantity = int(as_float)
            if loc is not None:
                result.entries.append((product, loc, quantity))

    if not result.entries and not result.errors:
        result.errors.append("登録できる持参数がありません")
    return result


def save_item_quantities(parsed: ParsedItemQuantity) -> int:
    """同じ 対象日×対象wk×メニュー×店 は上書き。保存件数を返す。"""
    # ItemQuantity の日付は文字列カラムに date を渡して "YYYY-MM-DD" で保存してきた（既存データと揃える）
    target_date = parsed.target_date.isoformat()
    target_week = parsed.target_week.isoformat()
    with transaction.atomic():
        for product, loc, quantity in parsed.entries:
            ItemQuantity.objects.update_or_create(
                target_date=target_date,
                target_week=target_week,
                product=product,
                sales_location=loc,
                defaults={"quantity": quantity},
            )
    return len(parsed.entries)


def is_business_day(d: date) -> bool:
    """土日・祝日・会社休日(Holiday)以外。"""
    if d.weekday() >= 5 or jpholiday.is_holiday(d):
        return False
    return not Holiday.objects.filter(date=d).exists()


def next_business_day(d: date) -> date:
    n = d + timedelta(days=1)
    while not is_business_day(n):
        n += timedelta(days=1)
    return n
